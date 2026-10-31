"""Native, editable reconstruction of Figure 2 (02_research_map) with zero dashes and zero defensive caveats."""
from pathlib import Path
import json, re, hashlib
import fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement

R = Path(__file__).resolve().parent
W, H = 1000, 1250
TEXT_ERRORS = []
C = {'teal': '168579', 'tl': 'ECF7F4', 'ink': '1E3045', 'muted': '6E7D8B', 'line': 'C9D1D9', 'white': 'FFFFFF', 'paper': 'FAF9F6', 'purple': '705C8D', 'pl': 'F0EAF7', 'blue': '346BC2', 'bl': 'EEF4FD', 'red': 'C75D50', 'rl': 'FCEFEA', 'green': '168579', 'gl': 'ECF7F4', 'gold': 'C68B36', 'yl': 'FDF2DD', 'gray': 'F4F3F4'}
fonts = {(b, i): fitz.Font(fontfile=str(R / 'fonts' / f"LiberationSans-{'BoldItalic' if b and i else 'Bold' if b else 'Italic' if i else 'Regular'}.ttf")) for b in [False, True] for i in [False, True]}

def rgb(c):
    return RGBColor.from_string(C.get(c, c))

def nofx(o):
    sp = o.element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
    if sp is not None:
        sp.append(OxmlElement('a:effectLst'))
    for e in o.element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):
        e.set('idx', '0')

class MapArt:
    def __init__(self, p, name):
        self.s = p.slides.add_slide(p.slide_layouts[6])
        self.name = name
        self.texts = []
        self.boxes = {}
        self.links = []
        self.s.background.fill.solid()
        self.s.background.fill.fore_color.rgb = rgb('white')

    def rect(self, n, x, y, w, h, fill='white', stroke='line', width=.8, rounding=.045, kind=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False):
        o = self.s.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
        o.name = n
        if fill is None:
            o.fill.background()
        else:
            o.fill.solid()
            o.fill.fore_color.rgb = rgb(fill)
        if stroke is None:
            o.line.fill.background()
        else:
            o.line.color.rgb = rgb(stroke)
            o.line.width = Pt(width)
        if dash:
            o.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            o.adjustments[0] = rounding
        nofx(o)
        self.boxes[n] = [x, y, w, h]
        return o

    def txt(self, n, x, y, w, h, t, size=16, bold=False, color='ink', align='left', italic=False, serif=False):
        font = fitz.Font(fontfile=str(R / 'fonts' / f"LiberationSerif-{'BoldItalic' if bold and italic else 'Bold' if bold else 'Italic' if italic else 'Regular'}.ttf")) if serif else fonts[(bold, italic)]
        lengths = [font.text_length(l, fontsize=size) for l in t.split('\n')]
        if max(lengths, default=0) > w - .2:
            TEXT_ERRORS.append((self.name, n, 'width', max(lengths), w, t))
        lh = size * 1.12
        if len(lengths) * lh > h + .2:
            TEXT_ERRORS.append((self.name, n, 'height', len(lengths) * lh, h, t))
        o = self.s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        o.name = n
        tf = o.text_frame
        tf.clear()
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf._txBody.bodyPr.set('anchorCtr', '0')
        for i, l in enumerate(t.split('\n')):
            q = tf.paragraphs[0] if not i else tf.add_paragraph()
            q.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
            q.space_before = q.space_after = Pt(0)
            q.line_spacing = Pt(lh)
            rr = q.add_run()
            rr.text = l
            rr.font.name = 'Liberation Serif' if serif else 'Liberation Sans'
            rr.font.size = Pt(size)
            rr.font.bold = bold
            rr.font.italic = italic
            rr.font.color.rgb = rgb(color)
        nofx(o)
        self.texts.append({'name': n, 'rect': [x, y, w, h], 'text': t, 'size': size, 'bold': bold, 'italic': italic, 'align': align})
        return o

    def line(self, n, a, b, color='ink', width=1):
        o = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(a[0]), Pt(a[1]), Pt(b[0]), Pt(b[1]))
        o.name = n
        o.line.color.rgb = rgb(color)
        o.line.width = Pt(width)
        nofx(o)
        return o

    def summary(self):
        return {'name': self.name, 'texts': self.texts, 'boxes': self.boxes, 'links': self.links, 'shape_count': len(self.s.shapes)}

D = json.loads((R / 'checks/map_data.json').read_text())
P = Presentation()
P.slide_width = Pt(W)
P.slide_height = Pt(H)
P.core_properties.title = 'Robust RAG research map (76 papers)'
a = MapArt(P, '02_research_map')
links = []

def hot(kind, target, x, y, w, h):
    links.append({'kind': kind, 'target': target, 'rect': [x, y, w, h]})

def t(n, x, y, w, h, s, size=18, bold=False, col='ink', align='left', serif=False):
    return a.txt(n, x, y, w, h, s, size, bold, col, align, serif=serif)

t('title', 55, 12, 930, 42, 'Robust RAG research map', 32, True, serif=True)
t('subtitle', 56, 56, 910, 23, '4 perspectives  /  20 strands  /  76 papers', 18, col='muted')
a.line('head_rule', (56, 87), (987, 87), 'line', .8)
for name, x, w, txt_lbl in [('family_head', 66, 151, 'PERSPECTIVE'), ('strand_head', 242, 221, 'RESEARCH STRAND'), ('paper_head', 485, 502, 'WORKS AND PUBLICATION YEARS')]:
    t(name, x, 98, w, 21, txt_lbl, 14.5, True, 'muted')

a.line('root_spine', (47, 256), (47, 1060), 'ink', 1.4)
a.rect('root_node', 9, 583, 38, 158, 'white', 'ink', 1.1, .08)
o = t('root_text', 11, 586, 160, 30, 'Robust RAG', 23, True, 'ink', 'center', True)
o.width = Pt(34)
o.height = Pt(152)
o.text_frame._txBody.bodyPr.set('vert', 'vert270')
o.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
hot('body', 'sec:definition', 9, 583, 38, 158)

colors = [('purple', 'pl'), ('blue', 'bl'), ('red', 'rl'), ('teal', 'tl')]
wrap = {
    'Filtering & compression': 'Filtering &\ncompression',
    'Memory--context conflict': 'Memory & context\nconflict',
    'Poisoning & backdoors': 'Poisoning &\nbackdoors',
    'Injection & persistence': 'Injection &\npersistence',
    'Sufficiency & answerability': 'Sufficiency &\nanswerability',
    'Uncertainty & calibration': 'Uncertainty &\ncalibration',
    'Verification & interaction': 'Verification &\ninteraction',
}
font = fonts[(False, False)]
for gi, g in enumerate(D['groups']):
    col, fill = colors[gi]
    y0 = 131 + 268 * gi
    cy = y0 + 125
    count = sum(len(row['works']) for row in g['rows'])
    a.line(f'root_branch{gi}', (47, cy), (66, cy), col, 1.4)
    a.rect(f'family{gi}', 66, cy - 53, 151, 106, fill, col, 1.0, .08)
    name = g['family'].replace('\\\\', '\n').replace('--', ' ').replace('–', ' ')
    t(f'family_name{gi}', 75, cy - 33, 133, 57, name, 23, True, col, 'center', True)
    t(f'family_sec{gi}', 73, cy + 31, 137, 19, f'§{gi + 4} · {count} papers', 14.5, col=col, align='center')
    hot('body', g['label'], 66, cy - 53, 151, 106)
    a.line(f'branch{gi}', (217, cy), (229, cy), col, 1.1)
    a.line(f'twig_spine{gi}', (229, y0 + 23), (229, y0 + 4 * 51 + 23), col, .9)
    for ri, row in enumerate(g['rows']):
        y = y0 + ri * 51
        n = f'{gi}_{ri}'
        a.line('strand_branch' + n, (229, y + 23), (242, y + 23), col, .85)
        a.rect('strand' + n, 242, y, 221, 46, 'white', col, .85, .045)
        title_str = wrap.get(row['title'], row['title']).replace('--', ' & ').replace('–', ' & ')
        lines = len(title_str.split('\n'))
        ts = 18.5
        th = lines * ts * 1.12
        t('strand_title' + n, 250, y + (46 - th) / 2 - 1, 205, th + 1, title_str, ts, True, col, 'center', True)
        hot('body', row['anchor'], 242, y, 221, 46)
        a.line('to_works' + n, (463, y + 23), (485, y + 23), col, .85)
        a.rect('work_band' + n, 485, y, 502, 46, fill if ri % 2 == 0 else 'white', col, .6, .025)
        a.line('column_divider' + n, (737, y + 7), (737, y + 39), 'line', .5)
        for wi, w in enumerate(row['works']):
            x = 494 + (wi % 2) * 251
            yy = y + 3 + (wi // 2) * 21
            alias = w['alias']
            m = re.search(r'\((\d{4})(?:\{\\natexlab\{([a-z])\}\})?', w['bibliography']['bibitem_label'])
            year = m[1] + (m[2] or '')
            sz = 17
            label = f'{alias} ({year})'
            total = font.text_length(label, fontsize=sz)
            assert total <= 238, (label, total)
            aw = font.text_length(alias, fontsize=sz)
            t('work_' + w['key'], x, yy, aw + .2, 21, alias, sz)
            t('year_' + w['key'], x + aw + 4, yy, 238 - aw - 4, 21, '(' + year + ')', sz, col='muted')
            hot('cite', w['key'], x, yy, total + 2, 21)

a.line('footer_rule', (56, 1210), (987, 1210), 'line', .8)
t('footer', 56, 1218, 930, 23, 'Systematic taxonomy of 76 representative works across four research perspectives and twenty mechanisms.', 16, col='muted')
assert not TEXT_ERRORS, json.dumps(TEXT_ERRORS, ensure_ascii=False, indent=2)
a.s.notes_slide.notes_text_frame.text = (
    'Image2 PNG design: image2_png/02_research_map.png\n'
    'SHA256=' + hashlib.sha256((R / 'image2_png/02_research_map.png').read_bytes()).hexdigest() + '\n'
    'Native editable paths, flat fills and text. Source: checks/map_data.json.'
)
P.save(R / 'pptx/02_research_map.pptx')
(R / 'checks/map_links.json').write_text(json.dumps(links, indent=2))
(R / 'checks/map_layout.json').write_text(json.dumps(a.summary(), indent=2))
print('NATIVE_PPT=' + str(R / 'pptx/02_research_map.pptx'))
print('NATIVE_SHAPES=' + str(len(a.s.shapes)))
print('PAPERS=76;STRANDS=20;HOTSPOTS=' + str(len(links)))
print('PREFLIGHT=PASS')
